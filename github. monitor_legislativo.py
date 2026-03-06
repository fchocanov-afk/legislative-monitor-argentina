#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MONITOR LEGISLATIVO - Cámara de Diputados Argentina
Fuente: https://www.diputados.gov.ar/proyectos/
Tema: Libertad de expresión e internet

Flujo:
1. POST al buscador con fecha 01/02/2026 hasta hoy → lista de proyectos
2. Por cada proyecto → POST a detalle_tp_adjunto → obtener URL del PDF
3. Descargar PDF → extraer texto
4. Claude lee texto → decide si es relevante + hace análisis completo
5. Guardar relevantes en Excel
"""

import json, re, sys, hashlib, logging, requests, time
import pandas as pd
from pathlib import Path
from datetime import datetime
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Directorios ────────────────────────────────────────────────────────────────
BASE_DIR   = Path.home() / "Desktop" / "monitoreo_legislativo"
CONFIG_F   = BASE_DIR / "config.json"
OUTPUT_DIR = BASE_DIR / "output"
LOG_DIR    = BASE_DIR / "logs"
PDF_DIR    = BASE_DIR / "pdfs"
HASHES_F   = BASE_DIR / "hashes_vistos.json"

for d in [OUTPUT_DIR, LOG_DIR, PDF_DIR]:
    d.mkdir(parents=True, exist_ok=True)

log_file = LOG_DIR / f"monitor_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(log_file, encoding="utf-8"),
        logging.StreamHandler()
    ]
)
log = logging.getLogger(__name__)

# ── URLs ───────────────────────────────────────────────────────────────────────
URL_BUSQUEDA  = "https://www.diputados.gov.ar/proyectos/resultado.html"
URL_DETALLE   = "https://www.hcdn.gob.ar/proyectos/detalle_tp_adjunto/index.html"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/145.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "es-AR,es;q=0.9,en;q=0.8",
    "Accept-Encoding": "gzip, deflate, br",
    "Connection": "keep-alive",
    "Referer": "https://www.diputados.gov.ar/proyectos/",
    "Origin": "https://www.diputados.gov.ar",
    # La cookie de sesión se carga desde config.json (campo "cookie")
    # No hardcodear aquí — ver config.example.json
}

# ── Config / hashes ────────────────────────────────────────────────────────────
def cargar_config():
    if not CONFIG_F.exists():
        log.error(f"No se encontro config.json en {CONFIG_F}")
        log.error("Copiá config.example.json como config.json y completá tus datos.")
        sys.exit(1)
    with open(CONFIG_F, encoding="utf-8") as f:
        return json.load(f)

def cargar_hashes():
    if HASHES_F.exists():
        with open(HASHES_F, encoding="utf-8") as f:
            return set(json.load(f))
    return set()

def guardar_hashes(h):
    with open(HASHES_F, "w", encoding="utf-8") as f:
        json.dump(list(h), f)

def hash_proyecto(expediente):
    return hashlib.md5(expediente.encode()).hexdigest()

# ── Scraping de resultados ─────────────────────────────────────────────────────
def buscar_proyectos(config):
    """
    Hace POST al buscador con fecha desde 01/02/2026 hasta hoy.
    Devuelve lista de dicts con los datos de cada proyecto.
    """
    fecha_desde = "01/02/2026"
    fecha_hasta = datetime.now().strftime("%d/%m/%Y")

    log.info(f"Buscando proyectos desde {fecha_desde} hasta {fecha_hasta}...")

    session = requests.Session()
    session.headers.update(HEADERS)

    # Inyectar cookie desde config si existe
    cookie = config.get("cookie", "")
    if cookie:
        session.headers["Cookie"] = cookie

    # Payload del formulario de búsqueda
    payload = {
        "strFechaInicio": fecha_desde,
        "strFechaFin":    fecha_hasta,
        "strCamara":      "D",  # Diputados
        "proyectos_por_pagina": "500",
    }

    proyectos  = []
    pagina     = 1
    total_pags = None

    while True:
        try:
            log.info(f"  Scrapeando página {pagina}...")

            payload_pag = dict(payload)
            if pagina > 1:
                payload_pag["pagina"] = str(pagina)

            r = session.post(URL_BUSQUEDA, data=payload_pag, timeout=60)

            if r.status_code != 200:
                log.warning(f"  Página {pagina}: status {r.status_code}")
                break

            soup = BeautifulSoup(r.text, "html.parser")

            if pagina == 1:
                titulo = soup.find("h3", class_="resultados-title")
                if titulo:
                    log.info(f"  {titulo.text.strip()}")
                ultima = soup.find("a", attrs={"aria-label": "Ultima"})
                if ultima:
                    href = ultima.get("href", "")
                    m = re.search(r"pagina=(\d+)", href)
                    if m:
                        total_pags = int(m.group(1))
                        log.info(f"  Total de páginas: {total_pags}")

            wrappers = soup.find_all("div", class_="resultado_wrapper")
            if not wrappers:
                log.info(f"  No hay más proyectos en página {pagina}")
                break

            for w in wrappers:
                p = parsear_wrapper(w)
                if p:
                    proyectos.append(p)

            log.info(f"  Página {pagina}: {len(wrappers)} proyectos (total: {len(proyectos)})")

            if total_pags and pagina >= total_pags:
                break
            siguiente = soup.find("a", attrs={"aria-label": "Siguiente"})
            if not siguiente:
                if total_pags and pagina < total_pags:
                    pass
                else:
                    break

            pagina += 1
            time.sleep(1.5)

        except Exception as e:
            log.error(f"  Error en página {pagina}: {e}")
            break

    log.info(f"Total proyectos encontrados: {len(proyectos)}")
    return proyectos, session


def parsear_wrapper(w):
    """Extrae datos de un bloque resultado_wrapper del HTML."""
    try:
        tipo = w.find("h4")
        tipo = tipo.text.strip() if tipo else ""

        meta = w.find("div", class_="dp-metadata")
        expediente = ""
        fecha      = ""
        publicado  = ""
        if meta:
            spans = meta.find_all("span")
            for s in spans:
                t = s.text.strip()
                if "Expediente Diputados:" in t:
                    strong = s.find("strong")
                    expediente = strong.text.strip() if strong else ""
                elif "Fecha:" in t:
                    strong = s.find("strong")
                    fecha = strong.text.strip() if strong else ""
                elif "Publicado en:" in t:
                    strong = s.find("strong")
                    publicado = strong.text.strip() if strong else ""

        titulo_div = w.find("div", class_="dp-texto")
        titulo = titulo_div.text.strip() if titulo_div else ""

        firmantes = []
        for tr in w.select("table tbody tr"):
            celdas = tr.find_all("td")
            if celdas:
                firmantes.append(celdas[0].text.strip())

        comisiones = []
        for h5 in w.find_all("h5"):
            if "COMISION" in h5.text.upper() or "COMISIÓN" in h5.text.upper():
                tabla = h5.find_next("table")
                if tabla:
                    for tr in tabla.select("tbody tr"):
                        comisiones.append(tr.text.strip())

        form = w.find("form", action=lambda a: a and "detalle_tp_adjunto" in a)
        id_interno = ""
        if form:
            inp = form.find("input", {"name": "id"})
            id_interno = inp["value"] if inp else ""

        if not expediente and not titulo:
            return None

        return {
            "expediente":  expediente,
            "fecha":       fecha,
            "publicado":   publicado,
            "tipo":        tipo,
            "titulo":      titulo,
            "firmantes":   ", ".join(firmantes),
            "comisiones":  ", ".join(comisiones),
            "id_interno":  id_interno,
            "link":        f"https://www.diputados.gov.ar/proyectos/resultado.html",
        }
    except Exception as e:
        log.warning(f"Error parseando wrapper: {e}")
        return None


# ── Obtener y descargar PDF ────────────────────────────────────────────────────
def obtener_url_pdf(id_interno, session):
    """Hace POST a detalle_tp_adjunto y extrae la URL del PDF."""
    try:
        r = session.post(
            URL_DETALLE,
            data={"id": id_interno},
            timeout=60,
            allow_redirects=True
        )
        if r.status_code != 200:
            return None

        soup = BeautifulSoup(r.text, "html.parser")

        match = re.search(r'https://rest\.hcdn\.gob\.ar/web/proyectos/\d+/adjuntos/\d+', r.text)
        if match:
            return match.group(0)

        for tag in soup.find_all(["iframe", "embed"]):
            src = tag.get("src", "")
            if "rest.hcdn" in src:
                return src

        return None
    except Exception as e:
        log.warning(f"Error obteniendo URL PDF para id {id_interno}: {e}")
        return None


def descargar_pdf(url_pdf, expediente, session):
    """Descarga el PDF y devuelve la ruta local."""
    try:
        nombre = re.sub(r'[^\w\-]', '_', expediente) + ".pdf"
        ruta   = PDF_DIR / nombre
        if ruta.exists():
            return ruta

        r = session.get(url_pdf, timeout=60, stream=True)
        if r.status_code == 200:
            with open(ruta, "wb") as f:
                for chunk in r.iter_content(chunk_size=8192):
                    f.write(chunk)
            return ruta
    except Exception as e:
        log.warning(f"Error descargando PDF {expediente}: {e}")
    return None


def extraer_texto_pdf(ruta_pdf):
    """Extrae texto del PDF."""
    if not ruta_pdf:
        return ""
    try:
        import fitz
        doc  = fitz.open(str(ruta_pdf))
        text = "\n".join(p.get_text() for p in doc)
        doc.close()
        return text[:15000]
    except Exception as e:
        log.warning(f"Error extrayendo texto PDF: {e}")
    return ""


# ── Análisis con Claude ────────────────────────────────────────────────────────
def analizar(proyecto, texto_pdf, config):
    """
    Claude lee el PDF y:
    1. Decide si el proyecto es relevante para libertad de expresión/internet
    2. Si es relevante, hace el análisis completo
    Devuelve dict con todos los campos, o {"relevante": "no"} si no es relevante.
    """
    import anthropic

    if texto_pdf.strip():
        contenido = texto_pdf
    else:
        contenido = f"""
Título: {proyecto.get('titulo', '')}
Tipo: {proyecto.get('tipo', '')}
Firmantes: {proyecto.get('firmantes', '')}
Comisiones: {proyecto.get('comisiones', '')}
"""

    prompt = f"""Sos experto en derecho constitucional y libertad de expresión en América Latina, con conocimiento del marco jurídico argentino (Constitución Nacional Art. 14 y 32, Ley 26.032, Convención Americana sobre DDHH) y los estándares de la CIDH.

Analizá el siguiente proyecto de ley del Congreso de Argentina.

PASO 1: Decidí si este proyecto es relevante para libertad de expresión, libertad de prensa, regulación de internet, plataformas digitales, datos personales, inteligencia artificial, intermediarios, o derechos digitales en general. Considerá también proyectos que parecen ser de otro tema (niñez, salud, educación) pero que regulan entornos digitales o plataformas.

PASO 2: Si es relevante, hacé el análisis completo.

Respondé SOLO con un JSON válido (sin texto extra, sin bloques markdown):

Si NO es relevante:
{{"relevante": "no"}}

Si SÍ es relevante:
{{
  "relevante": "si",
  "tipo": "tipo de instrumento (ley, resolución, declaración, pedido de informes, etc)",
  "extracto": "resumen de 2-3 oraciones del contenido",
  "oficialismo": "Si / No / No se puede determinar",
  "modifica_ley": "Si / No",
  "ley_que_modifica": "nombre o número de la ley, o vacío si no aplica",
  "estado_parlamentario": "estado conocido del trámite o vacío",
  "criminaliza_expresion": "Si / No / Parcialmente",
  "elimina_criminalizacion": "Si / No",
  "sancion_civil": "Si / No / Parcialmente",
  "elimina_sancion_civil": "Si / No",
  "barreras_expresion": "Si / No / Parcialmente",
  "regula_internet": "Si / No / Parcialmente",
  "distingue_online_offline": "Si / No / No aplica",
  "regula_intermediarios": "Si / No / Parcialmente",
  "apoyo_sociedad_civil": "Si / No / No se conoce",
  "limita_discurso": "Si / No / Parcialmente",
  "objetivo_legitimo": "Si / No / Parcialmente — breve explicación",
  "legalidad": "Si / No / Parcialmente — breve explicación",
  "necesidad": "Si / No / Parcialmente — breve explicación",
  "proporcionalidad": "Si / No / Parcialmente — breve explicación",
  "cumple_test": "Si / No / Parcialmente",
  "otros_objetivos_2": "segundo objetivo legítimo si existe, sino vacío",
  "otros_objetivos_3": "tercer objetivo legítimo si existe, sino vacío",
  "analisis_test_tripartito": "análisis del test tripartito CIDH en 3-5 oraciones",
  "recomendacion": "URGENTE / MONITOREAR / ARCHIVO",
  "justificacion_recomendacion": "justificación breve en 1-2 oraciones"
}}

Criterios de recomendación:
- URGENTE: restringe significativamente la libertad de expresión, criminaliza discurso, impone responsabilidad a intermediarios sin garantías, o no cumple el test tripartito CIDH
- MONITOREAR: tiene elementos ambiguos o podría tener impacto según cómo avance
- ARCHIVO: no tiene impacto relevante (declaraciones honoríficas, presupuestos positivos, etc)

PROYECTO:
Expediente: {proyecto.get('expediente', '')}
Título: {proyecto.get('titulo', '')}
Firmantes: {proyecto.get('firmantes', '')}
Comisiones: {proyecto.get('comisiones', '')}

TEXTO:
{contenido[:12000]}"""

    try:
        client  = anthropic.Anthropic(api_key=config["anthropic_api_key"])
        modelo  = config.get("llm_model", "claude-sonnet-4-20250514")
        message = client.messages.create(
            model=modelo,
            max_tokens=2000,
            messages=[{"role": "user", "content": prompt}]
        )
        respuesta = message.content[0].text.strip()
        respuesta = re.sub(r'^```json\s*', '', respuesta)
        respuesta = re.sub(r'\s*```$', '', respuesta)
        return json.loads(respuesta)
    except Exception as e:
        log.error(f"Error con Claude para {proyecto.get('expediente', '')}: {e}")
        return {"relevante": "no"}


# ── Guardar Excel ──────────────────────────────────────────────────────────────
def guardar_excel(filas, ruta_excel):
    if not filas:
        log.warning("No hay filas para guardar.")
        return

    df = pd.DataFrame(filas)

    columnas = [
        "Pais", "Anio", "N de expediente", "Origen", "Tipo", "Extracto", "Link",
        "Fecha de entrada", "Comision", "Cantidad de firmantes", "Partido politico",
        "Oficialismo", "Modifica ley vigente", "Ley que modifica",
        "Estado parlamentario", "Fecha de caducidad", "Aprobado en Camara origen",
        "Criminaliza la expresion", "Elimina criminalizacion",
        "Impone sancion civil", "Elimina sancion civil",
        "Impone barreras adicionales", "Regula contenido en Internet",
        "Distingue online de offline", "Regula intermediarios en internet",
        "Apoyo de sociedad civil", "Limita el discurso",
        "Objetivo legitimo", "Legalidad", "Necesidad", "Proporcionalidad",
        "Cumple con el test tripartito",
        "Otros objetivos 2", "Otros objetivos 3",
        "Analisis test tripartito",
        "Revision humana", "Notas",
        "Fecha procesado", "PDF local",
    ]

    for col in columnas:
        if col not in df.columns:
            df[col] = ""
    df = df[columnas]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Proyectos"

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )

    for ci, cn in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=ci, value=cn)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[1].height = 40

    fills = {
        "URGENTE":    PatternFill("solid", fgColor="FF9999"),
        "MONITOREAR": PatternFill("solid", fgColor="FFEB9C"),
        "ARCHIVO":    PatternFill("solid", fgColor="C6EFCE"),
    }
    col_rec = list(df.columns).index("Revision humana") + 1

    for ri, row in enumerate(df.itertuples(index=False), 2):
        rec_val = str(row[col_rec - 1]).upper() if row[col_rec - 1] else ""
        for ci, val in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=str(val) if val else "")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin
            if rec_val in fills:
                cell.fill = fills[rec_val]

    anchos = {
        "Extracto": 45, "Analisis test tripartito": 60,
        "Link": 35, "N de expediente": 18, "Fecha de entrada": 15,
        "Ley que modifica": 30, "Notas": 40,
    }
    for ci, cn in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(ci)].width = anchos.get(cn, 20)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(str(ruta_excel))
    log.info(f"Excel guardado: {ruta_excel}")


# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    log.info("=" * 60)
    log.info("MONITOR LEGISLATIVO - Cámara de Diputados Argentina")
    log.info("Tema: Libertad de expresión e internet")
    log.info("=" * 60)

    config  = cargar_config()
    hashes  = cargar_hashes()

    # 1. Buscar todos los proyectos del período
    proyectos, session = buscar_proyectos(config)
    if not proyectos:
        log.warning("No se encontraron proyectos.")
        return

    # 2. Filtrar los que ya fueron procesados
    nuevos = [(p, hash_proyecto(p["expediente"])) for p in proyectos
              if p["expediente"] and hash_proyecto(p["expediente"]) not in hashes]

    log.info(f"Proyectos nuevos a procesar: {len(nuevos)} de {len(proyectos)} total")

    filas = []
    for idx, (p, h) in enumerate(nuevos, 1):
        exp = p.get("expediente", "")
        log.info(f"[{idx}/{len(nuevos)}] {exp} — {p.get('titulo', '')[:70]}")

        # 3. Obtener URL del PDF y descargarlo
        url_pdf   = obtener_url_pdf(p["id_interno"], session) if p["id_interno"] else None
        ruta_pdf  = descargar_pdf(url_pdf, exp, session) if url_pdf else None
        texto_pdf = extraer_texto_pdf(ruta_pdf)

        if not url_pdf:
            log.warning(f"  Sin PDF para {exp}, usando solo título")

        # 4. Claude analiza
        a = analizar(p, texto_pdf, config)

        if a.get("relevante") == "no":
            log.info(f"  → No relevante, descartado")
            hashes.add(h)
            continue

        log.info(f"  → Relevante: {a.get('recomendacion', '?')}")

        anio_match = re.search(r'-(\d{4})$', exp)
        anio = anio_match.group(1) if anio_match else ""

        filas.append({
            "Pais":                              "Argentina",
            "Anio":                              anio,
            "N de expediente":                   exp,
            "Origen":                            p.get("firmantes", "").split(",")[0].strip(),
            "Tipo":                              a.get("tipo", p.get("tipo", "")),
            "Extracto":                          a.get("extracto", p.get("titulo", "")),
            "Link":                              url_pdf or p.get("link", ""),
            "Fecha de entrada":                  p.get("fecha", ""),
            "Comision":                          p.get("comisiones", ""),
            "Cantidad de firmantes":             len(p.get("firmantes", "").split(",")) if p.get("firmantes") else "",
            "Partido politico":                  "",
            "Oficialismo":                       a.get("oficialismo", ""),
            "Modifica ley vigente":              a.get("modifica_ley", ""),
            "Ley que modifica":                  a.get("ley_que_modifica", ""),
            "Estado parlamentario":              a.get("estado_parlamentario", ""),
            "Fecha de caducidad":                "",
            "Aprobado en Camara origen":         "",
            "Criminaliza la expresion":          a.get("criminaliza_expresion", ""),
            "Elimina criminalizacion":           a.get("elimina_criminalizacion", ""),
            "Impone sancion civil":              a.get("sancion_civil", ""),
            "Elimina sancion civil":             a.get("elimina_sancion_civil", ""),
            "Impone barreras adicionales":       a.get("barreras_expresion", ""),
            "Regula contenido en Internet":      a.get("regula_internet", ""),
            "Distingue online de offline":       a.get("distingue_online_offline", ""),
            "Regula intermediarios en internet": a.get("regula_intermediarios", ""),
            "Apoyo de sociedad civil":           a.get("apoyo_sociedad_civil", ""),
            "Limita el discurso":                a.get("limita_discurso", ""),
            "Objetivo legitimo":                 a.get("objetivo_legitimo", ""),
            "Legalidad":                         a.get("legalidad", ""),
            "Necesidad":                         a.get("necesidad", ""),
            "Proporcionalidad":                  a.get("proporcionalidad", ""),
            "Cumple con el test tripartito":     a.get("cumple_test", ""),
            "Otros objetivos 2":                 a.get("otros_objetivos_2", ""),
            "Otros objetivos 3":                 a.get("otros_objetivos_3", ""),
            "Analisis test tripartito":          a.get("analisis_test_tripartito", ""),
            "Revision humana":                   a.get("recomendacion", ""),
            "Notas":                             a.get("justificacion_recomendacion", ""),
            "Fecha procesado":                   datetime.now().strftime("%Y-%m-%d %H:%M"),
            "PDF local":                         str(ruta_pdf) if ruta_pdf else "",
        })
        hashes.add(h)
        time.sleep(0.5)

    # 5. Guardar Excel y hashes
    guardar_hashes(hashes)

    if filas:
        ts   = datetime.now().strftime("%Y%m%d_%H%M")
        ruta = OUTPUT_DIR / f"proyectos_arg_{ts}.xlsx"
        guardar_excel(filas, ruta)
        log.info(f"✅ {len(filas)} proyectos relevantes guardados.")
        log.info(f"📄 Excel: {ruta}")
    else:
        log.info("No se encontraron proyectos relevantes en este período.")

    log.info("Monitor finalizado.")


if __name__ == "__main__":
    main()
